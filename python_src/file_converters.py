"""File conversion utilities untuk Iqbal CV Bot"""
import os
import re
from openpyxl import load_workbook
import vobject

def create_vcf_entry(phone: str, name: str) -> str:
    """Create VCF contact entry"""
    # Clean phone number
    phone = re.sub(r'\D', '', phone)
    if not phone:
        return None
    
    return f"""BEGIN:VCARD
VERSION:3.0
FN:{name}
TEL;TYPE=CELL:+{phone}
END:VCARD
"""

def txt_to_vcf(txt_path: str, vcf_path: str, contact_name: str):
    """Convert TXT file to VCF"""
    try:
        with open(txt_path, 'r', encoding='utf-8') as f:
            content = f.read()
        
        # Extract phone numbers
        numbers = re.findall(r'[\d\+\-\(\) ]+', content)
        numbers = [re.sub(r'\D', '', n) for n in numbers]
        numbers = [n for n in numbers if n and len(n) >= 7]
        
        if not numbers:
            return False, "Tidak ada nomor yang ditemukan"
        
        # Create VCF
        vcf_content = ""
        for i, num in enumerate(numbers, 1):
            name = f"{contact_name} {i}"
            vcf_entry = create_vcf_entry(num, name)
            if vcf_entry:
                vcf_content += vcf_entry + "\n"
        
        with open(vcf_path, 'w', encoding='utf-8') as f:
            f.write(vcf_content)
        
        return True, f"Berhasil convert {len(numbers)} nomor"
    except Exception as e:
        return False, str(e)

def vcf_to_txt(vcf_path: str, txt_path: str):
    """Convert VCF file to TXT"""
    try:
        numbers = []
        
        with open(vcf_path, 'r', encoding='utf-8') as f:
            content = f.read()
        
        # Extract phone numbers from TEL fields
        tel_matches = re.findall(r'TEL[^:]*:(\+?\d+)', content)
        numbers = list(set(tel_matches))  # Remove duplicates
        
        if not numbers:
            return False, "Tidak ada nomor yang ditemukan"
        
        with open(txt_path, 'w', encoding='utf-8') as f:
            f.write('\n'.join(numbers))
        
        return True, f"Berhasil extract {len(numbers)} nomor"
    except Exception as e:
        return False, str(e)

def xls_to_vcf(xls_path: str, vcf_path: str, contact_name: str):
    """Convert XLSX file to VCF"""
    try:
        wb = load_workbook(xls_path)
        ws = wb.active
        
        numbers = []
        for row in ws.iter_rows(min_row=1):
            for cell in row:
                if cell.value:
                    num = re.sub(r'\D', '', str(cell.value))
                    if num and len(num) >= 7:
                        numbers.append(num)
        
        if not numbers:
            return False, "Tidak ada nomor di file XLSX"
        
        vcf_content = ""
        for i, num in enumerate(numbers, 1):
            name = f"{contact_name} {i}"
            vcf_entry = create_vcf_entry(num, name)
            if vcf_entry:
                vcf_content += vcf_entry + "\n"
        
        with open(vcf_path, 'w', encoding='utf-8') as f:
            f.write(vcf_content)
        
        return True, f"Berhasil convert {len(numbers)} nomor dari XLSX"
    except Exception as e:
        return False, str(e)

def clean_txt(txt_path: str) -> tuple:
    """Clean and sort TXT file (remove duplicates, trim whitespace)"""
    try:
        with open(txt_path, 'r', encoding='utf-8') as f:
            lines = f.readlines()
        
        # Clean and deduplicate
        numbers = set()
        for line in lines:
            num = re.sub(r'\D', '', line.strip())
            if num and len(num) >= 7:
                numbers.add(num)
        
        # Sort
        numbers = sorted(numbers)
        
        with open(txt_path, 'w', encoding='utf-8') as f:
            f.write('\n'.join(numbers))
        
        return True, f"Berhasil rapikan {len(numbers)} nomor"
    except Exception as e:
        return False, str(e)

def count_contacts(file_path: str) -> tuple:
    """Count contacts in VCF or TXT file"""
    try:
        if file_path.endswith('.vcf'):
            with open(file_path, 'r', encoding='utf-8') as f:
                content = f.read()
            count = len(re.findall(r'BEGIN:VCARD', content))
        elif file_path.endswith('.txt'):
            with open(file_path, 'r', encoding='utf-8') as f:
                lines = f.readlines()
            count = len([l for l in lines if l.strip()])
        else:
            return False, "Format tidak didukung"
        
        return True, f"Total: {count} kontak"
    except Exception as e:
        return False, str(e)

def split_vcf(vcf_path: str, output_dir: str, chunk_size: int = 100):
    """Split VCF file into multiple files"""
    try:
        with open(vcf_path, 'r', encoding='utf-8') as f:
            content = f.read()
        
        vcards = re.split(r'(?=BEGIN:VCARD)', content)
        vcards = [v.strip() for v in vcards if v.strip()]
        
        files = []
        for i in range(0, len(vcards), chunk_size):
            chunk = vcards[i:i+chunk_size]
            file_num = (i // chunk_size) + 1
            output_file = os.path.join(output_dir, f"split_{file_num}.vcf")
            
            with open(output_file, 'w', encoding='utf-8') as f:
                f.write('\n'.join(chunk))
            
            files.append(output_file)
        
        return True, files
    except Exception as e:
        return False, str(e)
